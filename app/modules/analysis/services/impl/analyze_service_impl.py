import base64
import json
import asyncio
import logging
from io import BytesIO
import openpyxl
import fitz
import docx
from typing import Any, Dict, List, Tuple
from fastapi import UploadFile

from app.modules.analysis.services.analyze_service import AnalyzeService
from app.modules.analysis.services.document_service import DocumentService
from app.modules.integration.extraction_engine import ExtractionEngine
from app.core.exceptions import AppBaseException
from app.utils.file_util import FileUtil
from app.dto.guia_aerea_dtos import GuiaAereaRequest
from app.dto.interviniente_dtos import IntervinienteRequest
from app.dto.confianza_extraccion_dtos import GuiaAereaConfianzaRequest
from app.core.tasks.document_tasks import process_document_validations

logger = logging.getLogger(__name__)

def process_file_content(content: bytes, filename: str) -> Tuple[List[str], str]:
    """
    CPU-bound task to convert PDF/Images to base64 OR Extract text from Excel.
    Executed in a thread pool.
    Returns: (list_of_base64_images, extracted_text_string)
    """
    base64_results = []
    text_result = ""

    if FileUtil.is_valid_pdf(content):
        try:
            doc = fitz.open(stream=BytesIO(content), filetype="pdf")
            for page in doc:
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
                img_bytes = pix.tobytes("jpeg")
                base64_results.append(base64.b64encode(img_bytes).decode("utf-8"))
            doc.close()
        except Exception as e:
            logger.error(f"Error processing PDF {filename}: {e}", exc_info=True)
            return [], ""

    elif FileUtil.is_valid_image(content):
        base64_results.append(base64.b64encode(content).decode("utf-8"))
    
    # Check for Excel
    elif FileUtil.is_valid_xlsx(content):
        try:
            wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
            text_parts = []
            for sheet in wb.worksheets:
                text_parts.append(f"--- HOJA: {sheet.title} ---")
                for row in sheet.iter_rows(values_only=True):
                    # Convert row to CSV-like string, filtering None
                    row_str = ",".join([str(c) for c in row if c is not None])
                    if row_str.strip():
                        text_parts.append(row_str)
            text_result = "\n".join(text_parts)
        except Exception as e:
            logger.error(f"Error processing Excel {filename}: {e}", exc_info=True)
            return [], ""

    # Check for Word
    elif FileUtil.is_valid_docx(content):
        try:
            doc = docx.Document(BytesIO(content))
            text_parts = []
            
            # Extract text from paragraphs
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        text_parts.append(" | ".join(row_text))
            
            text_result = "\n".join(text_parts)
        except Exception as e:
            logger.error(f"Error processing Word {filename}: {e}", exc_info=True)
            return [], ""

    return base64_results, text_result

class AnalyzeServiceImpl(AnalyzeService):

    def __init__(self, extraction_engine: ExtractionEngine):
        self.extraction_engine = extraction_engine

    async def upload(self, t: List[UploadFile]) -> List[Dict[str, Any]]:
        results = []
        loop = asyncio.get_running_loop()

        for file in t:
            content = await file.read()
            # Offload CPU-bound work to thread pool
            images, text = await loop.run_in_executor(None, process_file_content, content, file.filename)

            if not images and not text:
                logger.warning(f"File {file.filename} is corrupt or invalid.")
                raise AppBaseException(
                    message=f"El archivo {file.filename} está corrupto, vacío o tiene un formato no soportado."
                )
            
            file_extracted_data = []
            
            # Handle images (PDF/Images)
            if images:
                for idx, page_img in enumerate(images):
                    data = await self.extraction_engine.extract_data_from_image(page_img)
                    file_extracted_data.append({
                        "page": idx + 1,
                        "content": data
                    })
            # Handle text extraction (Excel/Word)
            elif text:
                 # Standardize to a single "page" for text-only documents in the legacy sync upload
                 text_b64 = base64.b64encode(text.encode('utf-8')).decode('utf-8')
                 # Use the generic document extraction for text
                 # Since upload is likely for single images, this part might need mapping if used
                 # For now, we reuse extraction_engine logic
                 results_list = await self.extraction_engine.extract_single_document(text_b64, "text/plain", 1, 1)
                 for idx, res in enumerate(results_list):
                     file_extracted_data.append({
                         "page": idx + 1,
                         "content": res.get("fields", {})
                     })

            results.append({
                "filename": file.filename,
                "total_pages": len(images) if images else 1,
                "data": file_extracted_data
            })

        return results


    async def upload_stream(self, files_data: List[Dict[str, Any]], document_service: DocumentService):
        all_docs_tasks = []
        total_files = len(files_data)
        
        # 1. Prepare all documents and calculate global indices
        current_global_page_index = 1
        for idx, file_data in enumerate(files_data):
            filename = file_data.get("filename")
            content = file_data.get("content")
            
            yield self._build_sse_event({"thinking": f"Preparando archivo {idx + 1}/{total_files}: {filename}...\n"})
            
            # Special case for Excel and Word - handle as text extraction
            if FileUtil.is_valid_xlsx(content) or FileUtil.is_valid_docx(content):
                 loop = asyncio.get_running_loop()
                 _, extracted_text = await loop.run_in_executor(None, process_file_content, content, filename)
                 if extracted_text:
                      text_b64 = base64.b64encode(extracted_text.encode('utf-8')).decode('utf-8')
                      task = self.extraction_engine.extract_single_document(
                          text_b64, 
                          "text/plain", 
                          1,
                          current_global_page_index
                      )
                      all_docs_tasks.append((task, filename, 1))
                      current_global_page_index += 1
                      continue

            # Standard case (PDF/Image) using demo's approach
            doc_prep = self.prepare_document_for_llm(content, filename)
            
            if not doc_prep:
                 yield self._build_sse_event({"thinking": f"[WARN] Archivo {filename} no soportado o vacío\n"})
                 continue

            page_count = doc_prep["page_count"]
            start_index = current_global_page_index
            
            # Schedule task for the entire document, passing page info
            task = self.extraction_engine.extract_single_document(
                doc_prep["base64"], 
                doc_prep["mime_type"], 
                page_count,
                start_index
            )
            all_docs_tasks.append((task, filename, page_count))
            
            # Increment global index for next file
            current_global_page_index += page_count
            logger.info(f"Added task for {filename} with {page_count} pages. Current global index: {current_global_page_index}")

        if not all_docs_tasks:
             logger.warning("No tasks were created for analysis.")
             yield self._build_sse_event({"thinking": "[ERROR] No se encontraron documentos válidos.\n"})
             return

        total_files_to_process = len(all_docs_tasks)
        yield self._build_sse_event({"thinking": f"Analizando {total_files_to_process} archivos en paralelo...\n"})

        # 2. Run in parallel
        try:
            completed_files_count = 0
            documents = []

            # Use demo's wrapper approach or similar to map results back to metadata
            # I'll use a safer task wrapping
            async def wrap_task(coro, fname, p_count):
                try:
                    return await coro, fname, p_count
                except Exception as e:
                    return [{"error": str(e)}], fname, p_count

            running_tasks = [wrap_task(t[0], t[1], t[2]) for t in all_docs_tasks]

            for future in asyncio.as_completed(running_tasks):
                results, fname, p_count = await future
                completed_files_count += 1
                
                yield self._build_sse_event({"thinking": f"Archivo '{fname}' ({p_count} pág) completado ({completed_files_count}/{total_files_to_process})\n"})
                
                if isinstance(results, list) and len(results) > 0 and results[0].get("error"):
                    err_msg = results[0].get("error")
                    logger.error(f"Error in task for {fname}: {err_msg}")
                    yield self._build_sse_event({"thinking": f"[ERROR] {fname}: {err_msg}\n"})
                else:
                    logger.info(f"Task for {fname} returned {len(results)} items.")
                    for page_result in results:
                        doc_obj = {
                            "document_index": page_result.get("document_index"),
                            "document_name": page_result.get("document_name") or f"{fname} - Pág {page_result.get('document_index')}",
                            "fileName": fname,
                            "fields": page_result.get("fields", {})
                        }
                        documents.append(doc_obj)

            # 3. Send Final JSON (sorted by document_index to keep order)
            documents.sort(key=lambda x: x.get("document_index", 0))
            logger.info(f"Finalizing analysis. Yielding {len(documents)} documents.")
            
            # Match demo exactly
            final_payload = json.dumps({"documents": documents})
            yield self._build_sse_event({"response": final_payload})
            
            # 4. AUTO-SAVE EXTRACTED DATA (Colleague's logic integration)
            if documents:
                yield self._build_sse_event({"thinking": "Registrando información en la base de datos...\n"})
                for doc_obj in documents:
                    try:
                        fields = doc_obj.get("fields", {})
                        if not fields or not fields.get("numero"):
                            logger.warning(f"Skipping auto-save for document '{doc_obj.get('fileName')}' - missing 'numero' field")
                            yield self._build_sse_event({"thinking": f"[SKIP] {doc_obj.get('fileName')}: No se detectó número de guía\\n"})
                            continue
                        
                        # Map fields to GuiaAereaRequest exactly as requested by colleague
                        # Extract intervinientes
                        intervinientes_data = fields.get("intervinientes", [])
                        intervinientes_req = []
                        for inter in intervinientes_data:
                            intervinientes_req.append(IntervinienteRequest(**inter))
                        
                        # Extract confianzas
                        confianzas_data = fields.get("confianzas", [])
                        confianzas_req = []
                        for conf in confianzas_data:
                            confianzas_req.append(GuiaAereaConfianzaRequest(**conf))

                        # Build the main request
                        # Note: We use all available fields from the LLM extraction
                        # For required fields, we provide defaults if missing
                        from datetime import datetime
                        
                        # Track missing critical fields for observation
                        missing_fields = []
                        
                        # Check critical fields
                        moneda = fields.get("monedaCodigo")
                        if not moneda:
                            missing_fields.append("monedaCodigo")
                            moneda = "USD"  # Default
                        
                        total_flete = fields.get("totalFlete")
                        # Validate that totalFlete is numeric
                        if total_flete:
                            try:
                                # Try to convert to float to validate
                                float(total_flete)
                            except (ValueError, TypeError):
                                # Non-numeric value like "AS ARRANGED"
                                logger.warning(f"Non-numeric totalFlete value: {total_flete}")
                                missing_fields.append("totalFlete")
                                total_flete = 0
                        else:
                            missing_fields.append("totalFlete")
                            total_flete = 0
                        
                        fecha_emision = fields.get("fechaEmision")
                        if not fecha_emision:
                            missing_fields.append("fechaEmision")
                            fecha_emision = datetime.now()
                        
                        origen = fields.get("origenCodigo")
                        if not origen:
                            missing_fields.append("origenCodigo")
                            origen = "N/A"
                        
                        destino = fields.get("destinoCodigo")
                        if not destino:
                            missing_fields.append("destinoCodigo")
                            destino = "N/A"
                        
                        cantidad = fields.get("cantidadPiezas")
                        # Validate that cantidadPiezas is numeric
                        if cantidad:
                            try:
                                int(cantidad)
                            except (ValueError, TypeError):
                                logger.warning(f"Non-numeric cantidadPiezas value: {cantidad}")
                                missing_fields.append("cantidadPiezas")
                                cantidad = 1
                        else:
                            missing_fields.append("cantidadPiezas")
                            cantidad = 1
                        
                        # Also validate other numeric fields and set to None if invalid
                        peso_bruto = fields.get("pesoBruto")
                        if peso_bruto:
                            try:
                                float(peso_bruto)
                            except (ValueError, TypeError):
                                logger.warning(f"Non-numeric pesoBruto value: {peso_bruto}, setting to None")
                                peso_bruto = None
                        
                        peso_cobrado = fields.get("pesoCobrado")
                        if peso_cobrado:
                            try:
                                float(peso_cobrado)
                            except (ValueError, TypeError):
                                logger.warning(f"Non-numeric pesoCobrado value: {peso_cobrado}, setting to None")
                                peso_cobrado = None
                        
                        guia_req = GuiaAereaRequest(
                            numero=fields.get("numero"),
                            fechaEmision=fecha_emision,
                            origenCodigo=origen,
                            destinoCodigo=destino,
                            aerolineaCodigo=fields.get("aerolineaCodigo"),
                            numeroVuelo=fields.get("numeroVuelo"),
                            fechaVuelo=fields.get("fechaVuelo"),
                            descripcionMercancia=fields.get("descripcionMercancia"),
                            cantidadPiezas=cantidad,
                            pesoBruto=peso_bruto,
                            pesoCobrado=peso_cobrado,
                            unidadPesoCodigo=fields.get("unidadPesoCodigo"),
                            totalFlete=total_flete,
                            monedaCodigo=moneda,
                            intervinientes=intervinientes_req,
                            confianzas=confianzas_req
                        )
                        
                        # If critical fields are missing, add observation
                        if missing_fields:
                            obs_msg = f"Campos críticos no detectados por la IA: {', '.join(missing_fields)}. Requiere revisión manual."
                            guia_req.observaciones = obs_msg
                            guia_req.estadoRegistroCodigo = "ESTGA002"  # OBSERVADO
                            logger.warning(f"AWB {guia_req.numero} marked as OBSERVADO due to missing fields: {missing_fields}")
                            yield self._build_sse_event({"thinking": f"[WARN] {doc_obj.get('fileName')}: Faltan campos críticos, marcada para revisión\\n"})
                        
                        # Persistent save via DocumentService
                        await document_service.saveOrUpdate(guia_req)
                        
                        # Trigger background validations exactly like DocumentFacade does
                        if guia_req.guiaAereaId:
                            process_document_validations.delay(guia_req.model_dump_json())
                            logger.info(f"Auto-saved AWB {guia_req.numero} (ID: {guia_req.guiaAereaId}) and triggered validations.")
                        
                    except Exception as save_err:
                        logger.error(f"Error auto-saving document {doc_obj.get('fileName')}: {save_err}")
                        yield self._build_sse_event({"thinking": f"[WARN] No se pudo auto-guardar {doc_obj.get('fileName')}: {str(save_err)}\n"})

            yield self._build_sse_event({"thinking": f"Análisis y registro finalizado exitosamente. {len(documents)} documentos procesados."})

        except Exception as e:
            logger.error(f"Critical Error during parallel extraction: {e}", exc_info=True)
            yield self._build_sse_event({"error": f"Error crítico: {str(e)}"})
            return


    def _build_sse_event(self, data: Dict[str, Any]) -> str:
        return f"data: {json.dumps(data)}\n\n"

    def prepare_document_for_llm(self, content: bytes, filename: str) -> Dict[str, Any]:
        """Returns base64, mime_type, and page_count for the document."""
        logger.info(f"Iniciando preparación de archivo: {filename} (Tamaño: {len(content)} bytes)")
        
        if FileUtil.is_valid_pdf(content):
            try:
                doc = fitz.open(stream=BytesIO(content), filetype="pdf")
                page_count = doc.page_count
                doc.close()
                logger.info(f"Archivo {filename} identificado como PDF con {page_count} páginas.")
                return {
                    "base64": base64.b64encode(content).decode("utf-8"),
                    "mime_type": "application/pdf",
                    "page_count": page_count
                }
            except Exception as e:
                logger.error(f"Error procesando PDF {filename}: {e}")
                
        if FileUtil.is_valid_image(content):
            mime_type = "image/jpeg"
            ext = filename.lower().split('.')[-1]
            if ext == "png":
                mime_type = "image/png"
            elif ext == "webp":
                mime_type = "image/webp"
            
            logger.info(f"Archivo {filename} identificado como IMAGEN con MIME: {mime_type}")
            return {
                "base64": base64.b64encode(content).decode("utf-8"),
                "mime_type": mime_type,
                "page_count": 1
            }
        
        logger.warning(f"Archivo {filename} no pudo ser identificado como PDF ni como Imagen.")
        return None
